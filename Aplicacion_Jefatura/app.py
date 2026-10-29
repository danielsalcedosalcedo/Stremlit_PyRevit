"""Dashboard ejecutivo de control de avance BIM, sin carga manual de archivos."""

from __future__ import annotations

import io
import re
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
import streamlit as st

ROOT_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT_DIR))

from logic.charts import (  # noqa: E402
    grafico_avance_por_item_plotly,
    grafico_comparativo_cantidades_plotly,
    grafico_curva_real_pyrevit_plotly,
    grafico_curva_s_diaria_plotly,
    grafico_curva_s_semanal_plotly,
)
from logic.loader import (  # noqa: E402
    load_avance_diario,
    load_avance_semanal,
    load_datos_sheet,
    load_excel_sheet_names,
    load_pyrevit_csv,
    load_pyrevit_json,
)
from logic.processor import (  # noqa: E402
    calcular_comparativo_cantidades_periodo,
    calcular_curva_s_diaria,
    calcular_curva_s_real_pyrevit,
    calcular_curva_s_semanal,
    detectar_unidad_pyrevit,
    extraer_cantidad_pyrevit,
    generar_reporte_diario_partidas,
)

st.set_page_config(
    page_title="Jefatura | Control BIM",
    page_icon="🔷",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Montserrat:wght@400;500;600;700;800&display=swap');
    :root {
        --ei-blue:#00A0E3; --ei-blue-dark:#0077B6; --ei-navy:#0A2A62;
        --panel:#F7F9FC; --border:#E6E9EE; --text:#2C3338;
        --sub:#6B7280; --green:#2E9E6B; --amber:#D4A017; --red:#C0392B;
    }
    html, body, [data-testid="stAppViewContainer"], .stApp {
        background:#fff !important; color:var(--text) !important;
        font-family:'Montserrat',system-ui,sans-serif !important;
    }
    .block-container {max-width:1320px; padding-top:.6rem; padding-bottom:2rem;}
    [data-testid="stSidebar"], [data-testid="stSidebarCollapsedControl"] {display:none;}
    [data-testid="stMetric"] {border:1px solid var(--border); border-radius:4px; padding:14px 16px; background:#fff;}
    [data-testid="stMetricLabel"] {color:var(--sub)!important; font-size:.72rem; font-weight:600; letter-spacing:.06em; text-transform:uppercase;}
    [data-testid="stMetricValue"] {color:var(--ei-navy)!important; font-weight:700;}
    [data-testid="stTabs"] button[aria-selected="true"] {color:var(--ei-blue)!important; border-bottom-color:var(--ei-blue)!important;}
    .jefatura-banner {padding:30px 34px; margin:8px 0 22px; color:#fff;
        background:linear-gradient(115deg,#0A2A62 0%,#1264A3 58%,#00A0E3 100%);}
    .jefatura-banner h1 {margin:0;color:#fff;font-size:1.8rem;font-weight:800;letter-spacing:.02em;}
    .jefatura-banner p {margin:8px 0 0;color:#EAF6FC;font-size:.9rem;}
    .section-title {font-size:.95rem;font-weight:700;color:var(--ei-navy);letter-spacing:.07em;
        text-transform:uppercase;border-bottom:2px solid var(--ei-blue);padding-bottom:8px;margin:1.4rem 0 .9rem;}
    .info-box {background:var(--panel);border:1px solid var(--border);border-left:3px solid var(--ei-blue);
        padding:12px 16px;margin:8px 0 14px;color:var(--sub);font-size:.85rem;}
    .source-pill {display:inline-block;padding:5px 10px;background:#EAF7FC;color:#0077B6;
        font-size:.72rem;font-weight:700;letter-spacing:.04em;text-transform:uppercase;}
    .footer {margin-top:2rem;padding:18px 4px;border-top:1px solid var(--border);color:#89929C;font-size:.72rem;}
    hr {border-color:var(--border)!important;}
    </style>
    """,
    unsafe_allow_html=True,
)

DEMOLICIONES_EXCEL_CANDIDATES = (
    ROOT_DIR / "data" / "programado" / "Demoliciones_Avance.xlsx",
    ROOT_DIR / "Demoliciones_Avance.xlsx",
)
HORMIGONES_EXCEL_CANDIDATES = (
    ROOT_DIR / "data" / "programado" / "Hormigones_Avance.xlsx",
    ROOT_DIR / "Hormigones_Avance.xlsx",
)
EEMM_EXCEL_CANDIDATES = (
    ROOT_DIR / "data" / "programado" / "EEMM_Avance.xlsx",
    ROOT_DIR / "EEMM_Avance.xlsx",
)
DEMOLICIONES_PYREVIT_CANDIDATES = (
    ROOT_DIR / "data" / "pyrevit_elementos_detallados.json",
    ROOT_DIR / "Aplicacion_pyRevit" / "data" / "pyrevit_elementos_detallados.json",
    ROOT_DIR / "data" / "pyrevit_elementos_detallados.csv",
)
HORMIGONES_PYREVIT_CANDIDATES = DEMOLICIONES_PYREVIT_CANDIDATES
EEMM_IFC_DIRECTORY = ROOT_DIR / "data" / "Modelos IFC"


def _first_existing(paths: tuple[Path, ...]) -> Path | None:
    return next((path for path in paths if path.is_file() and path.stat().st_size > 0), None)


@st.cache_data(show_spinner=False)
def _load_ifc_data(
    ifc_path: str,
    modified_at: float,
    include_geometry: bool,
    visible_ids: tuple[int, ...] = (),
) -> dict:
    """Carga propiedades IFC y cachea geometría hasta que cambia el archivo."""
    del modified_at
    from logic.ifc_viewer import load_ifc_model

    return load_ifc_model(
        ifc_path,
        include_geometry=include_geometry,
        element_ids=set(visible_ids) if include_geometry else None,
    )


@st.cache_data(show_spinner=False)
def _load_local_data(
    excel_path: str | None,
    excel_mtime: float,
    pyrevit_paths: tuple[str, ...],
    pyrevit_mtimes: tuple[float, ...],
    require_pyrevit: bool = True,
) -> dict:
    """Lee las fuentes fijas y renueva la caché cuando cambia su fecha de modificación."""
    del excel_mtime, pyrevit_mtimes  # Se conservan como parte de la clave de caché.
    result = {
        "df_datos": None,
        "avance_semanal": None,
        "avance_diario": None,
        "df_pyrevit": None,
        "pyrevit_path": None,
        "errores": [],
    }

    if excel_path:
        try:
            hojas = load_excel_sheet_names(excel_path)
            result["df_datos"] = load_datos_sheet(excel_path)
            if "Avance Semanal" in hojas:
                result["avance_semanal"] = load_avance_semanal(excel_path)
            if "Avance Diario" in hojas:
                result["avance_diario"] = load_avance_diario(excel_path)
        except Exception as exc:
            result["errores"].append(f"No se pudo leer el Excel programado: {exc}")

    if not require_pyrevit:
        return result

    if pyrevit_paths:
        errores_pyrevit = []
        try:
            for pyrevit_path in pyrevit_paths:
                try:
                    if pyrevit_path.lower().endswith(".csv"):
                        df_pyrevit = load_pyrevit_csv(pyrevit_path)
                    else:
                        df_pyrevit = load_pyrevit_json(pyrevit_path)

                    if df_pyrevit.empty:
                        errores_pyrevit.append(f"{pyrevit_path}: el archivo no contiene registros.")
                        continue
                    if "assembly_code" not in df_pyrevit.columns:
                        errores_pyrevit.append(f"{pyrevit_path}: falta la columna assembly_code.")
                        continue

                    result["df_pyrevit"] = df_pyrevit
                    result["pyrevit_path"] = pyrevit_path
                    break
                except Exception as exc:
                    errores_pyrevit.append(f"{pyrevit_path}: {exc}")
        except Exception as exc:
            errores_pyrevit.append(str(exc))

        if result["df_pyrevit"] is None:
            result["errores"].append(
                "No se pudo cargar una exportación BIM válida. " + " | ".join(errores_pyrevit)
            )
    else:
        result["errores"].append(
            "No se encontró ningún archivo de exportación PyRevit en las rutas del repositorio."
        )

    return result


def _filtered_as_of(df: pd.DataFrame, cutoff: date) -> pd.DataFrame:
    """Conserva todos los pendientes y solo los ejecutados hasta la fecha de corte."""
    if df.empty or "ejecutado" not in df.columns:
        return df.copy()

    result = df.copy()
    executed = result["ejecutado"].fillna(False).astype(bool)
    dates = pd.to_datetime(result.get("fecha"), format="mixed", dayfirst=True, errors="coerce")
    if "ei_fecha_ejecucion" in result.columns:
        fallback = pd.to_datetime(
            result["ei_fecha_ejecucion"], format="mixed", dayfirst=True, errors="coerce"
        )
        dates = dates.fillna(fallback)
    earliest = dates[executed].dropna().min()
    if pd.notna(earliest):
        dates = dates.fillna(earliest)

    result["_fecha_control"] = dates
    cutoff_ts = pd.Timestamp(cutoff)
    keep = (~executed) | (result["_fecha_control"].notna() & (result["_fecha_control"] <= cutoff_ts))
    return result.loc[keep].copy()


def _build_item_summary(
    df: pd.DataFrame,
    progress_as_of: pd.DataFrame | None = None,
    unit_override: str | None = None,
) -> pd.DataFrame:
    """Resume cantidades del alcance completo y avance opcional a una fecha de corte."""
    columns = [
        "ITEM",
        "assembly_code",
        "assembly_description",
        "Nombre de tarea",
        "Unidad",
        "Cantidad total",
        "Cantidad ejecutada",
        "pct_ejecutado",
        "Elementos",
        "Elementos ejecutados",
    ]
    if df.empty or "assembly_code" not in df.columns:
        return pd.DataFrame(columns=columns)

    rows = []
    for code, group in df.groupby("assembly_code", dropna=True):
        code_text = str(code).strip()
        if not code_text or code_text.lower() == "nan":
            continue

        description = ""
        if "assembly_description" in group.columns:
            modes = group["assembly_description"].dropna().astype(str).mode()
            description = modes.iloc[0] if not modes.empty else ""

        progress_group = group
        if progress_as_of is not None:
            progress_group = progress_as_of.loc[
                progress_as_of["assembly_code"].astype(str).str.strip() == code_text
            ]
        executed = (
            progress_group["ejecutado"].fillna(False).astype(bool)
            if "ejecutado" in progress_group
            else pd.Series(False, index=progress_group.index)
        )
        done = progress_group.loc[executed]
        unit = unit_override or detectar_unidad_pyrevit(group)
        total = float(extraer_cantidad_pyrevit(group, unit).sum())
        quantity_done = float(extraer_cantidad_pyrevit(done, unit).sum())
        pct = min(100.0, quantity_done / total * 100) if total > 0 else 0.0
        rows.append({
            "ITEM": code_text,
            "assembly_code": code_text,
            "assembly_description": description,
            "Nombre de tarea": description or code_text,
            "Unidad": unit,
            "Cantidad total": round(float(total), 2),
            "Cantidad ejecutada": round(float(quantity_done), 2),
            "pct_ejecutado": round(pct, 2),
            "Elementos": int(len(group)),
            "Elementos ejecutados": int(executed.sum()),
        })

    return pd.DataFrame(rows, columns=columns).sort_values(
        "pct_ejecutado", ascending=True
    )


def _build_weighted_global_curve(
    planned_data: dict,
    activity_choices: list[dict[str, str]],
    item_summary: pd.DataFrame,
    unit: str,
    frequency: str,
) -> tuple[pd.DataFrame, float, float]:
    """Calcula el programado global como cantidad BIM programada / cantidad BIM total."""
    if frequency == "W":
        curve = calcular_curva_s_semanal(planned_data)
    else:
        curve = calcular_curva_s_diaria(planned_data)
    if curve.empty:
        return curve, 0.0, 0.0

    unit_items = item_summary[item_summary["Unidad"] == unit]
    totals_by_code = dict(zip(unit_items["ITEM"], unit_items["Cantidad total"]))
    total_bim = float(unit_items["Cantidad total"].sum())
    if total_bim <= 0:
        return pd.DataFrame(), 0.0, 0.0

    activity_codes = {
        choice["activity"]: choice["code"]
        for choice in activity_choices
        if choice["activity"] in planned_data["actividades"]
    }
    weighted_percent = pd.Series(0.0, index=range(len(curve)))
    scheduled_bim = 0.0
    weighted_codes: set[str] = set()
    for activity_name, code_value in activity_codes.items():
        code = str(code_value).strip()
        weight = float(totals_by_code.get(code, 0.0))
        if weight <= 0 or code in weighted_codes:
            continue
        percentages = planned_data["actividades"][activity_name].get("acumulado", [])
        if len(percentages) != len(curve):
            continue
        weighted_percent = weighted_percent.add(
            pd.Series(percentages, dtype="float64") * weight,
            fill_value=0.0,
        )
        scheduled_bim += weight
        weighted_codes.add(code)

    curve["acumulado"] = (weighted_percent / total_bim).round(2)
    curve["parcial"] = curve["acumulado"].diff().fillna(curve["acumulado"])
    return curve, total_bim, scheduled_bim


def _normalize_task_name(value: object) -> str:
    """Normaliza nombres para enlazar las actividades programadas con Excel."""
    text = "" if value is None or pd.isna(value) else str(value).strip()
    text = re.sub(r"^\[[^\]]+\]\s*", "", text)
    return re.sub(r"\s+", " ", text).casefold()


def _activity_choices(planned_data: dict, df_datos: pd.DataFrame | None) -> list[dict[str, str]]:
    """Arma opciones de actividad mostrando ITEM y nombre de tarea."""
    activities = planned_data.get("actividades", {})
    rows = df_datos.to_dict("records") if df_datos is not None and not df_datos.empty else []
    valid_items = {
        str(row.get("ITEM", "")).strip()
        for row in rows
        if str(row.get("ITEM", "")).strip()
    }
    choices = []

    for activity_name, details in activities.items():
        scheduled_item = str(details.get("item", "")).strip()
        match = next(
            (
                row for row in rows
                if scheduled_item
                and str(row.get("ITEM", "")).strip() == scheduled_item
            ),
            None,
        )
        if match is None:
            activity_normalized = _normalize_task_name(activity_name)
            matches = [
                row for row in rows
                if _normalize_task_name(row.get("Nombre de tarea", "")) == activity_normalized
            ]
            if len(matches) == 1:
                match = matches[0]

        item_value = (
            match.get("ITEM", "")
            if match
            else scheduled_item
            if scheduled_item in valid_items
            or re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", scheduled_item)
            else ""
        )
        item_code = "" if pd.isna(item_value) else str(item_value).strip()
        task_value = match.get("Nombre de tarea", "") if match else activity_name
        task_name = "" if pd.isna(task_value) else str(task_value).strip()
        if not task_name:
            task_name = str(activity_name).strip()
        choices.append({
            "item": item_code or "Sin ITEM",
            "task": task_name,
            "activity": str(activity_name),
            "code": item_code,
        })

    return choices


def _build_excel_export(
    reports: dict[str, pd.DataFrame],
    codes_by_activity: dict[str, str] | None = None,
) -> bytes:
    """Construye un libro Excel con una hoja por partida."""
    output = io.BytesIO()
    used_names: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for report_key, report in reports.items():
            code = str(
                report.attrs.get("codigo_partida")
                or (codes_by_activity or {}).get(report_key)
                or report_key
            ).strip()
            sheet_name = re.sub(r"[\[\]:*?/\\]", "_", code).strip("' ")[:31]
            if not sheet_name:
                sheet_name = "Partida"
            base_name = sheet_name
            suffix = 2
            while sheet_name.casefold() in used_names:
                suffix_text = f"_{suffix}"
                sheet_name = f"{base_name[:31 - len(suffix_text)]}{suffix_text}"
                suffix += 1
            used_names.add(sheet_name.casefold())
            worksheet = writer.book.create_sheet(sheet_name)
            task_name = report.attrs.get("nombre_partida", report_key)
            worksheet.merge_cells("A1:E1")
            title = worksheet["A1"]
            title.value = f"{code} - {task_name}"
            title.font = Font(
                name="Calibri", size=14, bold=True, color="FFFFFF"
            )
            title.fill = PatternFill(
                fill_type="solid", fgColor="0A2A62"
            )
            title.alignment = Alignment(
                horizontal="left", vertical="center"
            )
            worksheet.row_dimensions[1].height = 26
            report.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=2,
            )
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A4"
            worksheet.auto_filter.ref = f"A3:E{worksheet.max_row}"
            worksheet.column_dimensions["A"].width = 15
            for column in ("B", "D"):
                worksheet.column_dimensions[column].width = 22
                for cell in worksheet[column][3:]:
                    cell.number_format = '0.00"%"'
            for column in ("C", "E"):
                worksheet.column_dimensions[column].width = 28
                for cell in worksheet[column][3:]:
                    cell.number_format = "#,##0.000"
            for cell in worksheet["A"][3:]:
                cell.number_format = "DD-MM-YYYY"
    return output.getvalue()


def _build_global_daily_reports(
    df_pyrevit: pd.DataFrame,
    item_summary: pd.DataFrame,
    activities: dict[str, dict],
    codes_by_activity: dict[str, str],
    fecha_inicio: pd.Timestamp,
    fecha_fin: pd.Timestamp,
    unit_labels: dict[str, str] | None = None,
) -> dict[str, pd.DataFrame]:
    """Genera avances diarios globales separados por unidad BIM."""
    start_date = pd.Timestamp(fecha_inicio).normalize()
    end_date = pd.Timestamp(fecha_fin).normalize()
    if start_date > end_date:
        raise ValueError("La fecha inicial del reporte no puede ser posterior a la fecha final.")

    dates = pd.date_range(start_date, end_date, freq="D")
    unit_labels = unit_labels or {
        "m2": "m²", "m3": "m³", "ml": "ml", "unid": "unitario"
    }
    if item_summary.empty:
        return {}

    totals_by_code = dict(zip(
        item_summary["assembly_code"].astype(str).str.strip(),
        item_summary["Cantidad total"],
    ))
    units_by_code = dict(zip(
        item_summary["assembly_code"].astype(str).str.strip(),
        item_summary["Unidad"],
    ))
    planned_by_unit = {
        unit: pd.Series(0.0, index=dates) for unit in unit_labels
    }
    scheduled_codes: set[str] = set()
    for activity, details in activities.items():
        code = str(codes_by_activity.get(activity) or details.get("item") or "").strip()
        if (
            not code
            or code in scheduled_codes
            or code not in totals_by_code
            or units_by_code[code] not in unit_labels
        ):
            continue
        scheduled_codes.add(code)

        activity_start = pd.to_datetime(details.get("inicio"), errors="coerce")
        duration_value = pd.to_numeric(details.get("dias"), errors="coerce")
        if pd.isna(activity_start) or pd.isna(duration_value) or duration_value <= 0:
            continue

        duration = int(round(float(duration_value)))
        if duration <= 0:
            continue
        workdays = pd.bdate_range(start=activity_start.normalize(), periods=duration)
        quantity_per_day = float(totals_by_code[code]) / duration
        unit = units_by_code[code]
        for workday in workdays:
            day = pd.Timestamp(workday).normalize()
            if start_date <= day <= end_date:
                planned_by_unit[unit].loc[day] += quantity_per_day

    actual_by_unit = {}
    for unit in unit_labels:
        actual_curve = calcular_curva_s_real_pyrevit(
            df_pyrevit,
            frecuencia="D",
            fecha_corte=end_date,
            unidad=unit,
        )
        actual_by_date = pd.Series(0.0, index=dates)
        for _, row in actual_curve.iterrows():
            day = pd.to_datetime(row["fecha"], errors="coerce")
            if pd.notna(day):
                day = day.normalize()
                if day in actual_by_date.index:
                    actual_by_date.loc[day] += float(
                        row["cantidad_ejecutada_periodo"]
                    )
        actual_by_unit[unit] = actual_by_date

    reports = {}
    for unit, unit_label in unit_labels.items():
        total_quantity = float(
            item_summary.loc[item_summary["Unidad"] == unit, "Cantidad total"].sum()
        )
        scheduled = planned_by_unit[unit]
        actual = actual_by_unit[unit]
        reports[unit] = pd.DataFrame({
            "Fecha": dates,
            "Avance programado (%)": (
                scheduled / total_quantity * 100 if total_quantity > 0 else 0.0
            ),
            "Avance programado (cantidad)": scheduled.to_numpy(),
            "Avance real (%)": (
                actual / total_quantity * 100 if total_quantity > 0 else 0.0
            ),
            "Avance real (cantidad)": actual.to_numpy(),
        })
        reports[unit].attrs["titulo"] = f"Global · Todas las partidas · {unit_label}"
    return reports


def _build_global_excel_export(reports: dict[str, pd.DataFrame]) -> bytes:
    """Construye un libro Excel global con una hoja por unidad."""
    output = io.BytesIO()
    unit_sheets = {
        "m2": "m2", "m3": "m3", "ml": "ml", "kg": "kg", "unid": "unitario"
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for unit, report in reports.items():
            sheet_name = unit_sheets[unit]
            worksheet = writer.book.create_sheet(sheet_name)
            worksheet.merge_cells("A1:E1")
            title = worksheet["A1"]
            title.value = report.attrs["titulo"]
            title.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            title.fill = PatternFill(fill_type="solid", fgColor="0A2A62")
            title.alignment = Alignment(horizontal="left", vertical="center")
            worksheet.row_dimensions[1].height = 26
            report.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=2,
            )
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A4"
            worksheet.auto_filter.ref = f"A3:E{worksheet.max_row}"
            worksheet.column_dimensions["A"].width = 15
            for column in ("B", "D"):
                worksheet.column_dimensions[column].width = 22
                for cell in worksheet[column][3:]:
                    cell.number_format = '0.00"%"'
            for column in ("C", "E"):
                worksheet.column_dimensions[column].width = 28
                for cell in worksheet[column][3:]:
                    cell.number_format = "#,##0.000"
            for cell in worksheet["A"][3:]:
                cell.number_format = "DD-MM-YYYY"
    return output.getvalue()


def _calcular_kpis_curva_s(
    planned_curve: pd.DataFrame,
    real_curve: pd.DataFrame,
    frequency: str,
    cutoff: date,
) -> dict[str, float | None]:
    """Calcula el porcentaje de avance programado, real y desviación/atraso a la fecha de corte."""
    pct_prog = None
    pct_real = None

    if not planned_curve.empty and "acumulado" in planned_curve.columns:
        if frequency == "W":
            from logic.charts import _current_week_label
            curr_lbl = _current_week_label(planned_curve, cutoff)
            if curr_lbl is not None and "semana_label" in planned_curve.columns:
                match = planned_curve[planned_curve["semana_label"].astype(str) == str(curr_lbl)]
                if not match.empty:
                    pct_prog = float(match.iloc[-1]["acumulado"])
            if pct_prog is None:
                pct_prog = float(planned_curve.iloc[-1]["acumulado"])
        else:
            fechas = pd.to_datetime(planned_curve.get("fecha", []), errors="coerce")
            cutoff_ts = pd.Timestamp(cutoff)
            valid = planned_curve[fechas <= cutoff_ts]
            if not valid.empty:
                pct_prog = float(valid.iloc[-1]["acumulado"])
            elif not planned_curve.empty:
                pct_prog = float(planned_curve.iloc[-1]["acumulado"])

    if not real_curve.empty and "pct_ejecutado_acum" in real_curve.columns:
        if "fecha" in real_curve.columns:
            fechas_r = pd.to_datetime(real_curve["fecha"], errors="coerce", dayfirst=True)
            cutoff_ts = pd.Timestamp(cutoff)
            valid_r = real_curve[fechas_r <= cutoff_ts]
            if not valid_r.empty:
                pct_real = float(valid_r.iloc[-1]["pct_ejecutado_acum"])
            else:
                pct_real = float(real_curve.iloc[-1]["pct_ejecutado_acum"])
        else:
            pct_real = float(real_curve.iloc[-1]["pct_ejecutado_acum"])

    desviacion = (pct_real - pct_prog) if (pct_prog is not None and pct_real is not None) else None
    return {
        "pct_prog": pct_prog,
        "pct_real": pct_real,
        "desviacion": desviacion,
    }


def _render_kpi_cards(cards: list[dict[str, str]]) -> None:
    """Renderiza tarjetas métricas ejecutivas limpias con color condicional y sin redundancia."""
    cards_html = []
    for card in cards:
        label = card.get("label", "")
        value = card.get("value", "")
        color = card.get("color", "#0A2A62")
        cards_html.append(
            f'<div style="flex:1 1 0;min-width:200px;border:1px solid #E6E9EE;border-radius:4px;padding:14px 16px;background:#fff;box-sizing:border-box;box-shadow:0 1px 2px rgba(0,0,0,0.03);">'
            f'<div style="color:#6B7280;font-size:0.72rem;font-weight:600;letter-spacing:0.06em;text-transform:uppercase;">{label}</div>'
            f'<div style="color:{color};font-size:1.75rem;font-weight:700;margin-top:6px;line-height:1.2;">{value}</div>'
            f'</div>'
        )
    html_code = f'<div style="display:flex;gap:14px;margin:8px 0 16px;flex-wrap:wrap;">{"".join(cards_html)}</div>'
    st.html(html_code)


def render_dashboard(
    excel_candidates: tuple[Path, ...],
    pyrevit_candidates: tuple[Path, ...],
    dashboard_key: str,
    forced_unit: str | None = None,
    ifc_directory: Path | None = None,
) -> None:
    cutoff = date.today()
    excel_file = _first_existing(excel_candidates)
    ifc_files = (
        tuple(sorted(ifc_directory.glob("*.ifc"), key=lambda path: path.name.casefold()))
        if ifc_directory and ifc_directory.is_dir()
        else ()
    )
    selected_ifc_files: list[Path] = []
    if ifc_directory:
        if not ifc_files:
            st.error(f"No se encontraron modelos IFC en {ifc_directory}.")
            return

        session_key = f"{dashboard_key}_ifc_models"
        if session_key not in st.session_state:
            # Selecciona el modelo más liviano por defecto para agilizar la carga inicial
            st.session_state[session_key] = [
                min(ifc_files, key=lambda path: path.stat().st_size)
            ]

        col_multiselect, col_btn1, col_btn2 = st.columns([6, 1.5, 1.5])
        with col_btn1:
            st.markdown("<div style='height: 28px;'></div>", unsafe_allow_html=True)
            if st.button("Marcar todos", key=f"{dashboard_key}_btn_all", use_container_width=True):
                st.session_state[session_key] = list(ifc_files)
                st.rerun()
        with col_btn2:
            st.markdown("<div style='height: 28px;'></div>", unsafe_allow_html=True)
            if st.button("Limpiar", key=f"{dashboard_key}_btn_clear", use_container_width=True):
                st.session_state[session_key] = []
                st.rerun()

        with col_multiselect:
            selected_ifc_files = st.multiselect(
                "Modelos IFC (puedes seleccionar uno o varios)",
                options=ifc_files,
                format_func=lambda path: path.name,
                key=session_key,
                help="Elige uno o más archivos IFC para consolidar avance y visualización 3D.",
            )

        if not selected_ifc_files:
            st.info("Selecciona al menos un archivo IFC para visualizar y analizar los datos de esta pestaña.")
            return

    pyrevit_files = tuple(
        path for path in pyrevit_candidates if path.is_file() and path.stat().st_size > 0
    )
    loaded = _load_local_data(
        str(excel_file) if excel_file else None,
        excel_file.stat().st_mtime if excel_file else 0.0,
        tuple(str(path) for path in pyrevit_files),
        tuple(path.stat().st_mtime for path in pyrevit_files),
        require_pyrevit=ifc_directory is None,
    )
    loaded_ifc_models: list[tuple[Path, dict]] = []
    if selected_ifc_files:
        for ifc_file in selected_ifc_files:
            try:
                model_data = _load_ifc_data(
                    str(ifc_file),
                    ifc_file.stat().st_mtime,
                    include_geometry=False,
                )
                loaded_ifc_models.append((ifc_file, model_data))
            except Exception as exc:
                st.warning(f"No fue posible leer el modelo IFC {ifc_file.name}: {exc}")

    all_records = []
    for ifc_file, model_data in loaded_ifc_models:
        for rec in model_data.get("records", []):
            if "archivo_ifc" not in rec:
                rec["archivo_ifc"] = ifc_file.name
            all_records.append(rec)

    ifc_records = (
        pd.DataFrame(
            all_records,
            columns=[
                "ifc_id", "global_id", "assembly_code", "assembly_description",
                "nombre", "kg", "cantidad_total", "unidad", "ejecutado",
                "ei_zona", "ei_subcontratista", "ei_semana",
                "ei_planificacion_obra", "ei_fecha_ejecucion", "fecha",
                "archivo_ifc",
            ],
        )
        if all_records
        else pd.DataFrame()
    )
    for error in loaded["errores"]:
        st.warning(error)

    if loaded["df_datos"] is None or "ITEM" not in loaded["df_datos"].columns:
        st.error("El Excel asignado a esta pestaña no contiene datos BIM con una columna ITEM.")
        return

    excel_codes = set(
        loaded["df_datos"]["ITEM"].dropna().astype(str).str.strip()
    )
    if not excel_codes and ifc_directory is None:
        st.info("El Excel asignado a esta pestaña no contiene códigos ITEM para vincular con BIM.")
        return

    raw_pyrevit = ifc_records if ifc_directory is not None else loaded["df_pyrevit"]
    if (raw_pyrevit is None or raw_pyrevit.empty) and ifc_directory is None:
        st.error("No hay datos BIM para mostrar. La aplicación no tiene carga manual: verifica la exportación automática en la carpeta data.")
        st.markdown(
            '<div class="info-box">Incluye una exportación PyRevit exclusiva para esta pestaña '
            'en una de las rutas indicadas. Cada pestaña consulta su propio archivo BIM.</div>',
            unsafe_allow_html=True,
        )
        with st.expander("Rutas BIM revisadas"):
            st.code("\n".join(str(path.relative_to(ROOT_DIR)) for path in pyrevit_candidates))
        return

    if raw_pyrevit is None:
        raw_pyrevit = pd.DataFrame()
    if "assembly_code" not in raw_pyrevit.columns and ifc_directory is None:
        st.error("La exportación encontrada no contiene el campo de Código de montaje (assembly_code).")
        return

    # Filtros de lectura del panel: zona y subcontratista.
    filter_cols = st.columns(2)
    with filter_cols[0]:
        zone_values = sorted(raw_pyrevit["ei_zona"].dropna().astype(str).unique()) if "ei_zona" in raw_pyrevit else []
        selected_zone = st.selectbox(
            "Zona", ["Todas"] + zone_values, key=f"{dashboard_key}_zone"
        )
    with filter_cols[1]:
        contractor_values = sorted(raw_pyrevit["ei_subcontratista"].dropna().astype(str).unique()) if "ei_subcontratista" in raw_pyrevit else []
        selected_contractor = st.selectbox(
            "Subcontratista",
            ["Todos"] + contractor_values,
            key=f"{dashboard_key}_contractor",
        )

    if not loaded_ifc_models and ifc_directory is None:
        raw_pyrevit = raw_pyrevit[
            raw_pyrevit["assembly_code"].astype(str).str.strip().isin(excel_codes)
        ].copy()
    elif excel_codes:
        raw_pyrevit = raw_pyrevit[
            raw_pyrevit["assembly_code"].astype(str).str.strip().isin(excel_codes)
        ].copy()
    else:
        raw_pyrevit = raw_pyrevit.iloc[0:0].copy()
    if raw_pyrevit.empty and not loaded_ifc_models and ifc_directory is None:
        st.info("La exportación BIM de esta pestaña no tiene códigos presentes en su Excel.")
        return
    if raw_pyrevit.empty and loaded_ifc_models:
        st.info(
            "Los IFC seleccionados aún no contienen Códigos de montaje vinculados al Excel. "
            "El visor se puede usar; agrega los parámetros BIM para habilitar los filtros y avances."
        )

    scope = raw_pyrevit.copy()
    if selected_zone != "Todas" and "ei_zona" in scope.columns:
        scope = scope[scope["ei_zona"].astype(str) == selected_zone]
    if selected_contractor != "Todos" and "ei_subcontratista" in scope.columns:
        scope = scope[scope["ei_subcontratista"].astype(str) == selected_contractor]

    as_of = _filtered_as_of(scope, cutoff)
    model_item_summary = _build_item_summary(scope, unit_override=forced_unit)
    item_summary = _build_item_summary(
        scope, progress_as_of=as_of, unit_override=forced_unit
    )
    st.markdown('<div class="section-title">Análisis ejecutivo</div>', unsafe_allow_html=True)
    dashboard_tabs = st.tabs(
        ["Curva de avance", "Partidas / códigos"]
        + (["Modelos IFC" if len(loaded_ifc_models) > 1 else "Modelo IFC"] if loaded_ifc_models else [])
    )
    tab_curve, tab_items = dashboard_tabs[:2]

    with tab_curve:
        frequency_label = st.radio(
            "Frecuencia",
            ["Semanal", "Diaria"],
            horizontal=True,
            key=f"{dashboard_key}_frequency",
        )
        frequency = "W" if frequency_label == "Semanal" else "D"
        planned_data = loaded["avance_semanal"] if frequency == "W" else loaded["avance_diario"]
        if planned_data:
            activity_choices = _activity_choices(planned_data, loaded["df_datos"])
        else:
            activity_choices = [
                {
                    "item": row["ITEM"],
                    "task": row["Nombre de tarea"],
                    "activity": "",
                    "code": row["ITEM"],
                }
                for _, row in item_summary.iterrows()
            ]
        chosen_activity = st.selectbox(
            "Seleccione partida o vista global",
            options=[None] + activity_choices,
            format_func=lambda option: (
                "Global · Todas las partidas"
                if option is None
                else f"{option['item']} · {option['task']}"
            ),
            key=f"{dashboard_key}_activity_{frequency}",
        )

        unit_labels = {"m2": "m²", "m3": "m³", "ml": "ml", "unid": "unidades"}
        if forced_unit:
            unit_labels = {
                forced_unit: (
                    "Kg" if forced_unit == "kg" else unit_labels[forced_unit]
                )
            }
        selected_unit = None
        if chosen_activity is None:
            available_units = list(unit_labels)
            selected_unit = st.selectbox(
                "Unidad de la curva global",
                options=available_units,
                index=(
                    available_units.index(forced_unit)
                    if forced_unit
                    else available_units.index("m2")
                ),
                format_func=lambda unit: unit_labels[unit],
                key=f"{dashboard_key}_global_unit_{frequency}",
                help="Cada curva global considera solo partidas de la unidad elegida.",
            )
            unit_total = (
                float(model_item_summary.loc[
                    model_item_summary["Unidad"] == selected_unit, "Cantidad total"
                ].sum())
                if not model_item_summary.empty
                else 0.0
            )
            if unit_total <= 0:
                st.warning(
                    f"Los datos BIM del filtro actual no contienen cantidades mayores que "
                    f"cero clasificadas como {unit_labels[selected_unit]}; no se puede "
                    "calcular esa curva."
                )

        curve_scope = scope
        if chosen_activity and chosen_activity["code"]:
            curve_scope = scope[
                scope["assembly_code"].astype(str).str.strip() == chosen_activity["code"]
            ]
        elif selected_unit:
            unit_codes = set(model_item_summary.loc[
                model_item_summary["Unidad"] == selected_unit, "assembly_code"
            ].astype(str))
            curve_scope = scope[
                scope["assembly_code"].astype(str).str.strip().isin(unit_codes)
            ]

        curve_unit = selected_unit
        if chosen_activity and chosen_activity["code"]:
            matching_summary = model_item_summary.loc[
                model_item_summary["assembly_code"] == chosen_activity["code"]
            ]
            if not matching_summary.empty:
                curve_unit = matching_summary.iloc[0]["Unidad"]

        real_curve = pd.DataFrame()
        if chosen_activity is not None or selected_unit:
            try:
                real_curve = calcular_curva_s_real_pyrevit(
                    curve_scope,
                    frecuencia=frequency,
                    fecha_corte=pd.Timestamp(cutoff),
                    unidad=curve_unit,
                )
            except Exception as exc:
                st.warning(f"No fue posible calcular la curva real: {exc}")

        if planned_data:
            activity = chosen_activity["activity"] if chosen_activity else None
            if chosen_activity is None and selected_unit:
                planned_curve, total_bim, scheduled_bim = _build_weighted_global_curve(
                    planned_data,
                    activity_choices,
                    model_item_summary,
                    selected_unit,
                    frequency,
                )
                if scheduled_bim < total_bim:
                    st.warning(
                        f"El Excel no tiene actividades vinculadas para "
                        f"{total_bim - scheduled_bim:,.2f} {unit_labels[selected_unit]} "
                        "del total BIM; esa cantidad queda incluida en el denominador "
                        "global, pero no aporta avance programado."
                    )
            elif chosen_activity is not None and frequency == "W":
                planned_curve = calcular_curva_s_semanal(planned_data, actividad=activity)
            elif chosen_activity is not None:
                planned_curve = calcular_curva_s_diaria(planned_data, actividad=activity)
            else:
                planned_curve = pd.DataFrame()

            if planned_curve.empty:
                st.info("No hay una curva programada disponible para la unidad seleccionada.")
            else:
                titulo_curva = (
                    f"Curva S · {unit_labels[selected_unit]} · Programado vs. real"
                    if chosen_activity is None and selected_unit
                    else "Curva S · Programado vs. real"
                    if frequency == "W"
                    else (
                        f"Curva diaria · {unit_labels[selected_unit]} · Programado vs. real"
                        if chosen_activity is None and selected_unit
                        else "Curva diaria · Programado vs. real"
                    )
                )

                # Indicador de % de atraso / adelanto a la fecha (sin valores redundantes y con color condicional)
                kpis = _calcular_kpis_curva_s(planned_curve, real_curve, frequency, cutoff)
                if kpis["pct_prog"] is not None or kpis["pct_real"] is not None:
                    prog_txt = f"{kpis['pct_prog']:.1f}%" if kpis["pct_prog"] is not None else "N/D"
                    real_txt = f"{kpis['pct_real']:.1f}%" if kpis["pct_real"] is not None else "0.0%"

                    if kpis["desviacion"] is not None:
                        desv = kpis["desviacion"]
                        if desv < -0.05:
                            estado_txt = f"Atraso de {abs(desv):.1f}%"
                            estado_color = "#C0392B"  # Rojo para atraso
                        elif desv > 0.05:
                            estado_txt = f"Adelanto de {desv:.1f}%"
                            estado_color = "#2E9E6B"  # Verde para adelanto
                        else:
                            estado_txt = "En plazo (0.0%)"
                            estado_color = "#0077B6"  # Azul para en plazo
                    else:
                        estado_txt = "Sin datos reales"
                        estado_color = "#6B7280"

                    _render_kpi_cards([
                        {"label": "Programado a la fecha", "value": prog_txt, "color": "#0A2A62"},
                        {"label": "Real ejecutado a la fecha", "value": real_txt, "color": "#0A2A62"},
                        {"label": "Estado / Atraso a la fecha", "value": estado_txt, "color": estado_color},
                    ])

                # Selector de motor de renderizado de la curva S
                col_chart_header, col_chart_ctrl = st.columns([7, 3])
                with col_chart_ctrl:
                    chart_engine = st.radio(
                        "Motor de gráfico",
                        ["Animado (JavaScript)", "Plotly"],
                        horizontal=True,
                        key=f"{dashboard_key}_engine_{frequency}",
                        help="Elige entre la vista animada de alto rendimiento con Apache ECharts o el visor clásico Plotly.",
                    )

                if chart_engine == "Animado (JavaScript)":
                    from logic.echarts_builder import build_curva_s_echarts_html
                    import streamlit.components.v1 as components

                    echarts_html = build_curva_s_echarts_html(
                        planned_curve,
                        real_curve,
                        frequency=frequency,
                        titulo=titulo_curva,
                        cutoff_date=cutoff,
                    )
                    components.html(echarts_html, height=500)
                else:
                    if frequency == "W":
                        chart = grafico_curva_s_semanal_plotly(
                            planned_curve,
                            titulo=titulo_curva,
                            df_real_pyrevit=real_curve,
                        )
                    else:
                        chart = grafico_curva_s_diaria_plotly(
                            planned_curve,
                            titulo=titulo_curva,
                            df_real_pyrevit=real_curve,
                        )
                    st.plotly_chart(
                        chart,
                        use_container_width=True,
                        key=f"{dashboard_key}_planned_curve_{frequency}",
                    )

            st.caption("La curva programada se obtiene del Excel fijo del proyecto; la real, de la exportación PyRevit.")
            scheduled_activities = {}
            if loaded["avance_diario"]:
                scheduled_activities = {
                    name: details.copy()
                    for name, details in loaded["avance_diario"]["actividades"].items()
                }
            activity_codes = {}
            for choice in activity_choices:
                name = choice["activity"]
                code = choice["code"]
                activity_codes[name] = code
                details = scheduled_activities.setdefault(name, {"item": code})
                if code and loaded["df_datos"] is not None and not loaded["df_datos"].empty:
                    schedule_row = loaded["df_datos"].loc[
                        loaded["df_datos"]["ITEM"].astype(str).str.strip() == code
                    ]
                    if not schedule_row.empty:
                        row = schedule_row.iloc[0]
                        if pd.isna(details.get("inicio")):
                            details["inicio"] = row.get("Comienzo")
                        if pd.isna(details.get("dias")):
                            details["dias"] = row.get("N° días hábiles (calc.)")

            if chosen_activity is None and selected_unit:
                selected_activities = {
                    choice["activity"] for choice in activity_choices
                    if choice["code"] in unit_codes
                }
                scheduled_activities = {
                    name: details for name, details in scheduled_activities.items()
                    if name in selected_activities
                }
                activity_codes = {
                    name: code for name, code in activity_codes.items()
                    if name in selected_activities
                }

            if not planned_curve.empty:
                comparison_unit = selected_unit
                if chosen_activity and chosen_activity["code"]:
                    matching_summary = model_item_summary.loc[
                        model_item_summary["assembly_code"]
                        == chosen_activity["code"]
                    ]
                    if not matching_summary.empty:
                        comparison_unit = matching_summary.iloc[0]["Unidad"]
                quantity_comparison = calcular_comparativo_cantidades_periodo(
                    planned_curve,
                    real_curve,
                    curve_scope,
                    scheduled_activities,
                    activity_codes,
                    frequency,
                    pd.Timestamp(cutoff),
                    actividad=chosen_activity["activity"] if chosen_activity else None,
                    unidad=comparison_unit,
                )
                st.markdown("#### Avance por período · Cantidad programada vs. real")

                # Métricas ejecutivas por cantidad (Programada, Real ejecutada y Estado/Atraso)
                unit_sym = unit_labels.get(comparison_unit, comparison_unit or "")
                if chosen_activity and chosen_activity["code"]:
                    code_str = str(chosen_activity["code"]).strip()
                    sub_model = model_item_summary.loc[model_item_summary["assembly_code"] == code_str]
                    sub_item = item_summary.loc[item_summary["assembly_code"] == code_str]
                else:
                    sub_model = model_item_summary.loc[model_item_summary["Unidad"] == comparison_unit]
                    sub_item = item_summary.loc[item_summary["Unidad"] == comparison_unit]

                total_cant = float(sub_model["Cantidad total"].sum()) if not sub_model.empty else 0.0
                real_cant = float(sub_item["Cantidad ejecutada"].sum()) if not sub_item.empty else 0.0
                prog_pct = kpis.get("pct_prog") if kpis else None
                prog_cant = total_cant * (prog_pct / 100.0) if (prog_pct is not None and total_cant > 0) else 0.0
                desv_cant = real_cant - prog_cant

                prog_cant_txt = f"{prog_cant:,.2f} {unit_sym}"
                real_cant_txt = f"{real_cant:,.2f} {unit_sym}"

                if prog_pct is not None and total_cant > 0:
                    if desv_cant < -0.01:
                        estado_cant_txt = f"Atraso de {abs(desv_cant):,.2f} {unit_sym}"
                        estado_cant_color = "#C0392B"  # Rojo para atraso
                    elif desv_cant > 0.01:
                        estado_cant_txt = f"Adelanto de {desv_cant:,.2f} {unit_sym}"
                        estado_cant_color = "#2E9E6B"  # Verde para adelanto
                    else:
                        estado_cant_txt = f"En plazo (0.00 {unit_sym})"
                        estado_cant_color = "#0077B6"  # Azul para en plazo
                else:
                    estado_cant_txt = f"{real_cant:,.2f} {unit_sym}"
                    estado_cant_color = "#0A2A62"

                _render_kpi_cards([
                    {"label": "Programado a la fecha", "value": prog_cant_txt, "color": "#0A2A62"},
                    {"label": "Real ejecutado a la fecha", "value": real_cant_txt, "color": "#0A2A62"},
                    {"label": "Estado / Atraso a la fecha", "value": estado_cant_txt, "color": estado_cant_color},
                ])

                if quantity_comparison.empty:
                    st.info("No hay cantidades comparables para el período y la unidad seleccionados.")
                else:
                    st.plotly_chart(
                        grafico_comparativo_cantidades_plotly(quantity_comparison),
                        use_container_width=True,
                        key=f"{dashboard_key}_period_comparison_{frequency}",
                    )
        else:
            if not excel_file:
                st.info(
                    f"La curva programada aparecerá cuando exista el Excel "
                    f"{excel_candidates[0].name}. La curva real BIM está disponible a continuación."
                )
            else:
                st.info("El Excel encontrado no incluye una hoja de avance para esta frecuencia.")
            st.plotly_chart(
                grafico_curva_real_pyrevit_plotly(
                    real_curve,
                    titulo=(
                        f"Curva real BIM · {unit_labels[selected_unit]} · "
                        f"{frequency_label.lower()} · al {cutoff.strftime('%d-%m-%Y')}"
                        if chosen_activity is None and selected_unit
                        else f"Curva real BIM · {frequency_label.lower()} · al "
                        f"{cutoff.strftime('%d-%m-%Y')}"
                    ),
                ),
                use_container_width=True,
                key=f"{dashboard_key}_real_curve_{frequency}",
            )

        st.markdown("#### Exportar evolución diaria de partidas")
        if loaded["avance_diario"]:
            export_choices = _activity_choices(
                loaded["avance_diario"], loaded["df_datos"]
            )
            export_codes = {
                choice["activity"]: choice["code"] for choice in export_choices
            }
            export_schedules = {
                name: details.copy()
                for name, details in loaded["avance_diario"]["actividades"].items()
            }
            for choice in export_choices:
                code = choice["code"]
                details = export_schedules.setdefault(
                    choice["activity"], {"item": code}
                )
                if not code or loaded["df_datos"] is None:
                    continue
                schedule_row = loaded["df_datos"].loc[
                    loaded["df_datos"]["ITEM"].astype(str).str.strip() == code
                ]
                if schedule_row.empty:
                    continue
                row = schedule_row.iloc[0]
                if pd.isna(details.get("inicio")):
                    details["inicio"] = row.get("Comienzo")
                if pd.isna(details.get("dias")):
                    details["dias"] = row.get("N° días hábiles (calc.)")

            try:
                export_start = pd.Timestamp("2026-08-24")
                export_end = pd.Timestamp(date.today())
                daily_reports = generar_reporte_diario_partidas(
                    scope,
                    export_schedules,
                    export_codes,
                    export_start,
                    export_end,
                )
                global_reports = _build_global_daily_reports(
                    scope,
                    model_item_summary,
                    export_schedules,
                    export_codes,
                    export_start,
                    export_end,
                    unit_labels=(
                        {forced_unit: unit_labels[forced_unit]}
                        if forced_unit
                        else None
                    ),
                )
                download_columns = st.columns(2)
                with download_columns[0]:
                    if daily_reports:
                        workbook = _build_excel_export(daily_reports, export_codes)
                        st.download_button(
                            "Descargar Excel con avance diario de partidas",
                            data=workbook,
                            file_name=(
                                f"avance_diario_partidas_2026-08-24_al_"
                                f"{date.today():%Y-%m-%d}.xlsx"
                            ),
                            mime=(
                                "application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sheet"
                            ),
                            key=f"{dashboard_key}_download_daily_progress",
                        )
                        st.caption(
                            f"El libro incluye {len(daily_reports)} partidas con cantidades BIM, "
                            "desde el 24-08-2026 hasta hoy. Los porcentajes corresponden al "
                            "avance diario respecto de la cantidad total de cada partida."
                        )
                    else:
                        st.info("No hay partidas con cantidad BIM y actividad programada para exportar.")
                with download_columns[1]:
                    if global_reports:
                        global_workbook = _build_global_excel_export(global_reports)
                        st.download_button(
                            "Descargar Excel Global",
                            data=global_workbook,
                            file_name=(
                                f"avance_global_todas_las_partidas_2026-08-24_al_"
                                f"{date.today():%Y-%m-%d}.xlsx"
                            ),
                            mime=(
                                "application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sheet"
                            ),
                            key=f"{dashboard_key}_download_global_daily_progress",
                        )
                        st.caption(
                            "Incluye hojas separadas para "
                            f"{', '.join(unit_labels.values())}. "
                            "Los porcentajes diarios se calculan sobre la cantidad BIM "
                            "total de cada unidad."
                        )
                    else:
                        st.info("No hay cantidades BIM disponibles para generar el avance global.")
            except Exception as exc:
                st.error(f"No fue posible generar el Excel de avance diario: {exc}")
        else:
            st.info("La exportación requiere la hoja Avance Diario del Excel programado.")

    with tab_items:
        if item_summary.empty:
            st.info("No hay datos agrupables por Código de montaje.")
        else:
            selected_codes = st.multiselect(
                "Filtrar códigos de montaje",
                options=item_summary["ITEM"].tolist(),
                default=[],
                help="Sin selección se muestran todos los códigos.",
                key=f"{dashboard_key}_item_codes",
            )
            visible_items = item_summary[item_summary["ITEM"].isin(selected_codes)] if selected_codes else item_summary
            table = visible_items[[
                "ITEM", "assembly_description", "Unidad", "Cantidad total", "Cantidad ejecutada",
                "pct_ejecutado", "Elementos", "Elementos ejecutados",
            ]].rename(columns={"assembly_description": "Descripción", "pct_ejecutado": "% Avance"})
            table["Unidad"] = table["Unidad"].replace({"kg": "Kg"})
            st.dataframe(
                table,
                use_container_width=True,
                hide_index=True,
                column_config={
                    "% Avance": st.column_config.ProgressColumn("% Avance", min_value=0, max_value=100, format="%.1f%%"),
                    "Cantidad total": st.column_config.NumberColumn("Cantidad total", format="%.2f"),
                    "Cantidad ejecutada": st.column_config.NumberColumn("Cantidad ejecutada", format="%.2f"),
                },
            )
            st.plotly_chart(
                grafico_avance_por_item_plotly(visible_items, titulo="Avance por código de montaje"),
                use_container_width=True,
                key=f"{dashboard_key}_item_progress",
            )

    if loaded_ifc_models:
        with dashboard_tabs[2]:
            viewer_scope = ifc_records.copy()
            if selected_zone != "Todas" and "ei_zona" in viewer_scope.columns:
                viewer_scope = viewer_scope[
                    viewer_scope["ei_zona"].astype(str) == selected_zone
                ]
            if (
                selected_contractor != "Todos"
                and "ei_subcontratista" in viewer_scope.columns
            ):
                viewer_scope = viewer_scope[
                    viewer_scope["ei_subcontratista"].astype(str)
                    == selected_contractor
                ]
            if excel_codes:
                linked_elements = viewer_scope[
                    viewer_scope["assembly_code"].astype(str).str.strip().isin(excel_codes)
                ]
                if not linked_elements.empty:
                    viewer_scope = linked_elements

            st.markdown("#### Visor IFC · elementos ejecutados en azul")
            st.caption(
                "El visor respeta los filtros de zona y subcontratista seleccionados arriba."
            )
            if len(loaded_ifc_models) > 1:
                module_options = [path.name for path, _ in loaded_ifc_models]
                selected_3d_modules = st.multiselect(
                    "Modelos activos en la vista 3D",
                    options=module_options,
                    default=module_options,
                    help="Permite prender o apagar modelos individuales en el visor 3D.",
                    key=f"{dashboard_key}_ifc_3d_modules",
                )
                if selected_3d_modules:
                    viewer_scope = viewer_scope[
                        viewer_scope["archivo_ifc"].isin(selected_3d_modules)
                    ]
                else:
                    st.info("Selecciona al menos un modelo para visualizar en 3D.")

            available_codes = sorted(
                code for code in viewer_scope["assembly_code"].dropna()
                .astype(str).str.strip().unique()
                if code
            )
            visible_codes = st.multiselect(
                "Filtrar códigos de montaje en el modelo",
                options=available_codes,
                default=[],
                help="Sin selección se muestran todos los elementos del alcance filtrado.",
                key=f"{dashboard_key}_ifc_codes",
            )
            if visible_codes:
                viewer_scope = viewer_scope[
                    viewer_scope["assembly_code"].astype(str).str.strip().isin(visible_codes)
                ]

            executed_count = int(viewer_scope["ejecutado"].fillna(False).astype(bool).sum())
            total_weight = pd.to_numeric(
                viewer_scope["kg"], errors="coerce"
            ).fillna(0).sum()
            metric_columns = st.columns(3)
            metric_columns[0].metric("Elementos visibles", f"{len(viewer_scope):,}")
            metric_columns[1].metric("Elementos ejecutados", f"{executed_count:,}")
            metric_columns[2].metric("Peso del alcance", f"{total_weight:,.2f} Kg")
            missing_weights = int(viewer_scope["kg"].isna().sum())
            if missing_weights:
                st.warning(
                    f"{missing_weights} elementos del alcance no tienen un peso "
                    "numérico en la propiedad Weigth."
                )
            if not available_codes:
                st.info(
                    "Los modelos IFC seleccionados aún no tienen el parámetro Código de montaje. "
                    "Los elementos se visualizarán en gris hasta agregar los parámetros de avance."
                )
            st.caption(
                "Los elementos marcados como ejecutados se resaltan en azul; "
                "los restantes se muestran en gris. Usa el mouse para rotar y acercar el modelo."
            )
            if st.checkbox(
                "Cargar geometría 3D del IFC",
                value=False,
                key=f"{dashboard_key}_load_ifc_geometry",
            ):
                try:
                    models_with_geometry = []
                    total_missing_geometry = 0
                    active_files = [
                        path for path, _ in loaded_ifc_models
                        if not viewer_scope.empty and (viewer_scope["archivo_ifc"] == path.name).any()
                    ]
                    with st.spinner(f"Procesando geometría 3D de {len(active_files)} modelo(s) IFC..."):
                        for ifc_file in active_files:
                            file_scope = viewer_scope[viewer_scope["archivo_ifc"] == ifc_file.name]
                            file_visible_ids = set(
                                pd.to_numeric(file_scope["ifc_id"], errors="coerce")
                                .dropna().astype(int)
                            )
                            if not file_visible_ids:
                                continue
                            model_with_geom = _load_ifc_data(
                                str(ifc_file),
                                ifc_file.stat().st_mtime,
                                include_geometry=True,
                                visible_ids=tuple(sorted(file_visible_ids)),
                            )
                            models_with_geometry.append((model_with_geom, file_visible_ids))
                            if model_with_geom.get("elements_without_geometry"):
                                total_missing_geometry += model_with_geom["elements_without_geometry"]

                    from logic.ifc_viewer import build_ifc_figure

                    if models_with_geometry:
                        figure = build_ifc_figure(models_with_geometry)
                        if figure.data:
                            st.plotly_chart(
                                figure,
                                use_container_width=True,
                                key=f"{dashboard_key}_ifc_viewer",
                            )
                        else:
                            st.info(
                                "No hay geometría visible para la combinación actual de filtros."
                            )
                        if total_missing_geometry:
                            st.warning(
                                f"{total_missing_geometry} elementos IFC no tienen geometría renderizable."
                            )
                    else:
                        st.info("No hay elementos visibles para cargar geometría.")
                except Exception as exc:
                    st.error(f"No fue posible visualizar los modelos IFC: {exc}")
            else:
                st.info(
                    "Activa la casilla para procesar y mostrar la geometría 3D de los modelos seleccionados."
                )

st.markdown(
    """<div class="jefatura-banner">
        <div class="source-pill">Dashboard ejecutivo · Control BIM</div>
        <h1>Boulevard Gastronomico Mall Plaza el Trebol</h1>
    </div>""",
    unsafe_allow_html=True,
)

tab_demoliciones, tab_hormigones, tab_eemm = st.tabs(
    ["Demoliciones", "Hormigones", "EEMM"]
)
with tab_demoliciones:
    render_dashboard(
        DEMOLICIONES_EXCEL_CANDIDATES,
        DEMOLICIONES_PYREVIT_CANDIDATES,
        "demoliciones",
    )
with tab_hormigones:
    render_dashboard(
        HORMIGONES_EXCEL_CANDIDATES,
        HORMIGONES_PYREVIT_CANDIDATES,
        "hormigones",
    )
with tab_eemm:
    render_dashboard(
        EEMM_EXCEL_CANDIDATES,
        (),
        "eemm",
        forced_unit="kg",
        ifc_directory=EEMM_IFC_DIRECTORY,
    )

st.markdown(
    '<div class="footer">© Desarrollado por Daniel Salcedo Salcedo 2026-Coordinador BIM · Datos consultados automáticamente desde las fuentes locales del proyecto.</div>',
    unsafe_allow_html=True,
)
